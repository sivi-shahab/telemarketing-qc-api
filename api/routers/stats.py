import io
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment
from sqlalchemy.orm import Session

from api.dependencies import (
    get_current_user,
    get_evaluation_detail_user,
    get_db,
)
from qc_core.sales_lookup import (
    agent_ids_for_agent,
    agent_ids_for_am,
    agent_ids_for_hierarchy_filter,
    agent_ids_for_tl,
    hierarchy_filter_options,
)
from api.schemas.result import (
    DailyStatsResponse,
    NamaIbuKandungResponse,
    ResultListItem,
    ResultListResponse,
    StatsResponse,
    TicketDeleteResponse,
)
from api.permissions import (
    ADMIN_TICKET_DELETE,
    ERROR_CODE_APPEAL,
    ERROR_CODE_REVIEW_SPQ,
    MANUAL_STATUS_REVIEW_SPQ,
    RESULTS_EXPORT_TICKETS,
    RESULTS_EXPORT_VERIFICATION,
    RESULTS_FILTER_QC_SIDE,
    RESULTS_STATUS_REASON_FULL,
    STATS_RISK_BASE,
    STATS_RISK_SYSTEM_NEW,
    SCOPE_QC_ASSIGNED,
    SCOPE_QC_SUPPORT_OWN,
    SCOPE_SALES_AGENT,
    SCOPE_SALES_AM,
    SCOPE_SALES_TL,
    STATS_FAILURE_REASON,
    STATS_QC_PERFORMANCE,
    is_sales_scope,
)
from api.rbac import data_scope_for, effective_campaigns_for, has_perm, require
from api.qc_scope import ensure_can_view_result, scoped_customer_ids
from qc_core.db import crud
from qc_core.compliance.badwords import badword_fail_reason, badword_rows, has_badword
from qc_core.compliance.error_codes import (
    _appeal_kind,
    added_appeals_only,
    apply_added_score_appeals,
    inject_added_rows,
    merge_dynamic_verification_rows,
    apply_approved_appeals,
    apply_approved_card_holder_appeals,
    apply_approved_cashline_appeals,
    apply_approved_critical_compliance_appeals,
    annotate_critical_compliance_reasons,
    appeals_that_flip,
    approved_appeals_only,
    CARD_HOLDER_DYNAMIC_FIELDS,
    CARD_HOLDER_STATIC_SCORECARD,
    build_error_code_table,
    card_holder_two_match_satisfied,
    code_for_cashline_field,
    effective_appeal_status,
    not_fulfilled_reason,
    apply_cashline_document_status,
    apply_static_document_status,
    normalize_dynamic_verification,
    normalize_static_verification,
    static_consistency_failures,
)
from qc_core.compliance.documents import (
    DOCUMENT_TYPES,
    card_holder_bands_apply,
    required_doc_requirements,
    format_similarity,
)
from qc_core.compliance.reference_data import get_credit_limit, get_customer_info, npwp_required_by_limit
from qc_core.compliance.scoring import (
    max_score,
    no_product_interest,
    phase3_score,
    score_bomb_items,
    scorecard_score,
)
from qc_core.compliance.stats_aggregate import (
    data_gap_map,
    data_gap_reasons,
    doc_sla_enabled,
    _doc_sla_expired,
    category_label,
    effective_manual_status,
    manual_review_state,
    manual_status_of,
    _missing_docs_map,
    _normalized_json,
    document_status_map,
    _parse_ymd,
    ai_status_for_result,
    ai_status_map,
    compute_ai_status_timeseries,
    compute_stats_snapshot,
)

router = APIRouter(dependencies=[Depends(get_current_user)])

# Export column order (per tambahan.md).
_EXPORT_COLUMNS = ["id", "category", "item_code", "description", "score", "evidence", "ticket_id"]
_ERROR_COLUMNS = ["id", "sumber", "error_code", "item_code", "details_error", "reason", "evidence", "ticket_id"]
_RINGKASAN_COLUMNS = ["keterangan", "perubahan", "hasil"]

# TMS change-flag key -> human label surfaced as an Upload Document trigger. Order
# here is the display order in the button tooltip.
_CHANGE_LABELS = {
    "kantor": "Perubahan Alamat Kantor",
    "rumah": "Perubahan Alamat Rumah",
    "npwp": "Perubahan NPWP",
    "nik": "Perubahan NIK",
}

# Urutan slot dokumen (sama dengan crud.allowed_doc_types_from_flags).
_DOC_TYPE_ORDER = ["ktp", "npwp", "kk", "cover_buku_tabungan"]


def _doc_label(doc_type: str) -> str:
    return DOCUMENT_TYPES.get(doc_type, {}).get("label", doc_type.upper())


def _document_requirements(flags: dict, raw_result_json, bands_apply: bool, npwp_by_limit: bool) -> list[dict]:
    """Dokumen yang diminta tiket ini BESERTA alasannya, satu entri per pasangan
    (dokumen, alasan): ``{"doc_type", "doc_label", "reason"}``.

    Sebelumnya alasan ("Perubahan NPWP", "Limit >= 50 jt", band similarity) dan
    daftar dokumen dikirim sebagai dua daftar terpisah, sehingga kolom Document
    tidak bisa menyebut alasan mana milik dokumen mana. Kolom itu sekarang
    merangkainya menjadi "Perlu Dokumen NPWP karena Perubahan NPWP".
    """
    out: list[dict] = []

    def add(doc_type: str, reason: str) -> None:
        if any(o["doc_type"] == doc_type and o["reason"] == reason for o in out):
            return
        out.append({"doc_type": doc_type, "doc_label": _doc_label(doc_type), "reason": reason})

    # 1) Perubahan data TMS -> dokumen sesuai crud.CHANGE_DOC_TYPES.
    for key, label in _CHANGE_LABELS.items():
        if flags.get(key):
            for doc_type in crud.CHANGE_DOC_TYPES.get(key, []):
                add(doc_type, label)
    # 2) Band similarity card holder (Fase #5): nama ibu kandung 80-89% -> KK,
    #    tanggal lahir 87,5-99% -> KTP.
    # Sejak 24 Agustus 2026 daftar ini juga memuat cover buku tabungan dari
    # cashline_data_verification (nama_pemilik_rekening MISMATCH), lewat
    # required_doc_requirements — sebelumnya kewajibannya hanya berupa kalimat di
    # "reason" dan tidak pernah muncul sebagai dokumen yang diminta.
    if bands_apply:
        for req in required_doc_requirements(raw_result_json):
            sim = req.get("similarity")
            why = f"{req['label']} mirip {format_similarity(sim)}%" if sim is not None else req["label"]
            add(req["doc_type"], why)
    # 3) Limit pencairan >= Rp 50 juta -> NPWP.
    if npwp_by_limit:
        add("npwp", "Limit >= 50 jt")
    return sorted(out, key=lambda r: _DOC_TYPE_ORDER.index(r["doc_type"])
                  if r["doc_type"] in _DOC_TYPE_ORDER else len(_DOC_TYPE_ORDER))


def _scoped_customer_ids(db: Session, current_user) -> Optional[list]:
    """The customer/ticket ids a scoped role may see, else ``None`` (no scoping):

    - ``qc``: only the tickets a Team Leader QC has ASSIGNED to them (manual QC
      assignment). No assignment => empty list => sees nothing.
    - ``area_manager``: tickets of every agent under their ``NIP AM`` (== username) &
      DEDICATED=cashline — i.e. their whole area (all team leaders + agents below).
    - ``team_leader``: tickets of agents under their ``NIP TL`` (== username) &
      DEDICATED=cashline — i.e. their whole team.
    - ``sales_agent`` (individual agent): tickets of the agent whose ``NIP BARU``
      (== username) & DEDICATED=cashline — i.e. their own tickets.
    - ``qc_support``: hanya tiket complaint yang di-upload QC Support (himpunan
      terisolasi yang sama dengan filter daftar Results-nya).
    - ``spq_head`` / ``team_leader_qc`` / ``telesales_head`` / system: ``None`` =>
      see everything.

    Implementasinya tinggal di ``api.qc_scope`` supaya router Results, Statistics dan
    Transcripts memakai definisi cakupan yang sama persis; alias ini dipertahankan
    karena namanya sudah dipakai di banyak tempat pada modul ini.
    """
    return scoped_customer_ids(db, current_user)


def _customer_id_from_files(source_files) -> Optional[str]:
    """Derive the customer/ticket id (prefix before the first ``_``) from the
    first source filename, e.g. ``181001bWvi_20260521150215.pdf`` -> ``181001bWvi``."""
    if not source_files:
        return None
    first = source_files[0]
    if not isinstance(first, str) or not first:
        return None
    return first.split("_", 1)[0]


def _evaluation_dict(result_json) -> Optional[dict]:
    """Return the LLM ``evaluation`` object from a stored result_json, or None."""
    if not isinstance(result_json, dict):
        return None
    evaluation = result_json.get("evaluation")
    return evaluation if isinstance(evaluation, dict) else None


def _numeric_or_none(value) -> Optional[float]:
    """Return ``value`` as a number (int when whole), or None if not numeric."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return int(value) if value == int(value) else value


def _to_num(value) -> Optional[float]:
    """Coerce ``value`` (number or numeric string) to a number (int when whole),
    or None if not numeric."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        n = float(value)
    else:
        try:
            n = float(value)
        except (TypeError, ValueError):
            return None
    if n != n:  # NaN
        return None
    return int(n) if n == int(n) else n


# Skor dipinjam dari ``compliance/scoring.py`` — SATU implementasi untuk seluruh
# sistem. Sampai 28 Agustus 2026 berkas ini menyimpan salinannya sendiri, dan
# salinan itu tertinggal saat ZERO-SCORE RULE ditambahkan di modul aslinya: tabel
# Results dan export XLSX memberi PASS kepada tiket yang oleh halaman detail sudah
# dinyatakan FAIL. Alias dipertahankan supaya pemanggil di bawah tidak perlu diubah.
_max_score = max_score
_scorecard_score = scorecard_score
_phase3_score = phase3_score
_score_bomb_items = score_bomb_items
_no_product_interest = no_product_interest


def _manual_status(qc_req, ai_status) -> Optional[str]:
    """The "Manual Status" column: the verdict on the ticket as it stands for a HUMAN
    — 'PASS' (Qualified) / 'FAIL' (Not Qualified) / 'PENDING'.

    Sejak tiket pertama kali dibuat nilainya MENGIKUTI AI Status; begitu QC / TL QC /
    SPQ Head menetapkan vonisnya DAN vonis itu disetujui, vonis itulah yang dipakai —
    dan sejak saat itu AI Status ikut terkunci ke nilai yang sama (aturan 7 Agustus
    2026, lihat ``_result_ai_status``). Jadi kedua kolom tidak pernah berbeda; yang
    membedakan "dinilai mesin" dari "dinilai manusia" adalah ``manual_status_by_human``.

    Keadaan alur kerjanya (menunggu review / usulan ditolak / kurang dokumen) adalah
    sumbu terpisah — lihat ``manual_review_state``."""
    return effective_manual_status(qc_req, ai_status)


def _appeal_summary(appeals: list) -> Optional[dict]:
    """Summarize Error Code appeals ("banding") for one result, for the Results
    list "Banding Review" column: latest-per-row status counts + full history.
    QC uses this to see SPQ Head's review outcomes; SPQ Head uses ``pending`` to
    spot which IDs are awaiting review."""
    if not appeals:
        return None
    # Latest appeal per (error_code, item_code) — appeals arrive oldest-first.
    latest = {}
    for a in appeals:
        latest[(a.error_code, a.item_code)] = a
    counts = {"pending": 0, "approved": 0, "rejected": 0}
    tl_pending = 0   # awaiting Team Leader QC decision
    spq_pending = 0  # escalated by TL QC, awaiting SPQ Head
    for a in latest.values():
        eff = effective_appeal_status(a)
        counts[eff] = counts.get(eff, 0) + 1
        tl = getattr(a, "tl_qc_status", "pending")
        if tl == "pending":
            tl_pending += 1
        elif tl == "escalated" and a.approval_status == "pending":
            spq_pending += 1
    history = [
        {
            "error_code": a.error_code,
            "item_code": a.item_code,
            "approval_status": a.approval_status,
            "tl_qc_status": getattr(a, "tl_qc_status", "pending"),
            "tl_qc_username": getattr(a, "tl_qc_username", None),
            # Dibutuhkan agar Results bisa merender timeline banding yang SAMA dengan
            # tabel Error Code (lihat dashboard/src/utils/appealTimeline.js).
            "tl_qc_reviewed_at": a.tl_qc_reviewed_at.isoformat() if getattr(a, "tl_qc_reviewed_at", None) else None,
            "appeal_kind": getattr(a, "appeal_kind", "remove"),
            "origin": getattr(a, "origin", "qc"),
            "qc_reason": a.qc_reason,
            "tl_qc_comment": getattr(a, "tl_qc_comment", None),
            "review_comment": getattr(a, "review_comment", None),
            "requested_by_username": a.requested_by_username,
            "reviewed_by_username": a.reviewed_by_username,
            "reviewed_at": a.reviewed_at.isoformat() if a.reviewed_at else None,
            "requested_at": a.requested_at.isoformat() if a.requested_at else None,
        }
        for a in appeals
    ]
    return {
        "total": len(latest),
        "pending": counts["pending"],
        "approved": counts["approved"],
        "rejected": counts["rejected"],
        "tl_pending": tl_pending,
        "spq_pending": spq_pending,
        "history": history,
    }


@router.get("/stats/qc_performance")
def qc_performance(
    campaign: Optional[str] = Query(None, description="Batasi ke tiket satu campaign"),
    db: Session = Depends(get_db),
    current_user=Depends(require(STATS_QC_PERFORMANCE)),
):
    """Per-QC assigned / approved / approve-rate table, shown beneath the
    Hierarki Failure Rate tree. Restricted to Team Leader QC, SPQ Head and admin —
    the roles that manage the QC division. Dibatasi ke campaign yang menjadi cakupan
    pemanggil."""
    return crud.qc_performance_rows(db, campaign,
                                    effective_campaigns_for(db, current_user))


@router.get("/results/hierarchy_options")
def results_hierarchy_options(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """AM / TL / TLO options for the Results hierarchy filter, scoped to the
    caller's role. Empty lists mean the caller has no subordinates and the UI
    should hide the filter entirely.

    Team Leader QC oversees the QC team, not the sales org, so it gets QC / QC
    Support dropdowns (``qc_users`` / ``qc_support_users``) INSTEAD of AM/TL/TLO."""
    if has_perm(db, current_user, RESULTS_FILTER_QC_SIDE):
        qc = crud.qc_side_filter_options(db)
        return {
            "area_managers": [], "team_leaders": [], "agents": [],
            "qc_users": qc["qc_users"], "qc_support_users": qc["qc_support_users"],
        }
    opts = hierarchy_filter_options(db, data_scope_for(db, current_user),
                                    getattr(current_user, "username", None),
                                    effective_campaigns_for(db, current_user))
    opts["qc_users"] = []
    opts["qc_support_users"] = []
    return opts


def _resolve_filtered_results(
    db,
    current_user,
    *,
    status=None,
    campaign=None,
    ticket_id=None,
    ai_status=None,
    manual_status=None,
    am_nip=None,
    tl_nip=None,
    agent_nip=None,
    qc_username=None,
    qc_support_username=None,
    date_start=None,
    date_end=None,
    banding_pending=False,
    manual_status_pending=False,
    page=1,
    limit=20,
):
    """Baris hasil yang cocok dengan filter menu Results — TANPA pengayaan.

    Dipakai bersama oleh ``/list_results`` (yang lalu mengayakan tiap baris) dan
    Reprocess All (yang hanya butuh daftar tiketnya). Pemisahan ini bukan sekadar
    kerapian: filter ``ai_status``/``manual_status`` diturunkan di Python, bukan di
    SQL, jadi satu-satunya cara agar jumlah tiket yang DILIHAT Admin di modal
    konfirmasi sama dengan yang DIKERJAKAN adalah memakai jalur pemilihan yang sama
    persis. Dua salinan logika filter pasti menyimpang — seperti salinan
    ``_max_score``/``_scorecard_score`` yang dulu tertinggal di berkas ini.

    Mengembalikan ``(rows, total)``. Seluruh pengaman cakupan (campaign milik role,
    ``data_scope``, isolasi QC Support) ada DI SINI, jadi pemanggil mana pun
    mewarisinya dan tidak ada jalan pintas RBAC.
    """
    d_start = _parse_ymd(date_start)
    d_end = _parse_ymd(date_end)
    # Pembatasan CAMPAIGN milik role: mempersempit, tidak pernah memperlebar. Role
    # tanpa daftar campaign (bawaan semua role sistem, termasuk qc) tidak terpengaruh.
    # Kalau user memfilter campaign di luar cakupannya, hasilnya sengaja kosong
    # ketimbang diam-diam melebar ke campaign lain.
    role_campaigns = effective_campaigns_for(db, current_user)
    if role_campaigns is not None and campaign:
        allowed = {c.strip().casefold() for c in role_campaigns}
        if campaign.strip().casefold() not in allowed:
            return [], 0
    # Sales Agent (TL) & QC (agent) are scoped to their tickets; other roles: all.
    scoped_cids = _scoped_customer_ids(db, current_user)
    # Hierarchy dropdown (AM / TL / TLO). Narrows WITHIN the role scope above —
    # it never widens it, so an Area Manager passing another AM's NIP still only
    # sees their own tickets (the intersection is empty).
    filter_uids = agent_ids_for_hierarchy_filter(db, am_nip, tl_nip, agent_nip,
                                                 effective_campaigns_for(db, current_user))
    if filter_uids is not None:
        filter_cids = crud.customer_ids_for_agent_ids(db, list(filter_uids))
        scoped_cids = (
            list(set(scoped_cids) & set(filter_cids)) if scoped_cids is not None
            else list(filter_cids)
        )
    # QC Support is a standalone, isolated result set: it sees ONLY its own uploads
    # (all complaint tickets); every other role EXCLUDES QC Support's uploads.
    _iso = ({"uploaded_by_role": "qc_support"}
            if data_scope_for(db, current_user) == SCOPE_QC_SUPPORT_OWN
            else {"exclude_uploaded_by_role": "qc_support"})
    # Team Leader QC filters by QC team member instead of the sales hierarchy:
    #  - "Semua QC"        -> narrow to tickets ASSIGNED to that QC;
    #  - "Semua QC Support" -> show that QC Support's uploads (overrides the default
    #    qc_support exclusion so those complaint tickets become visible).
    if has_perm(db, current_user, RESULTS_FILTER_QC_SIDE):
        if (qc_username or "").strip():
            qc_cids = crud.assigned_ticket_ids_for_qc(db, qc_username)
            scoped_cids = (
                list(set(scoped_cids) & set(qc_cids)) if scoped_cids is not None
                else list(qc_cids)
            )
        if (qc_support_username or "").strip():
            _iso = {"uploaded_by_username": qc_support_username.strip()}
    ai_filter = ai_status.strip().upper() if isinstance(ai_status, str) else None
    manual_filter = manual_status.strip().upper() if isinstance(manual_status, str) else None
    _has_status_filter = (ai_filter in ("PASS", "FAIL", "PENDING")
                          or manual_filter in ("PASS", "FAIL", "PENDING"))

    def _apply_status_filters(rows: list) -> list:
        """Saring ``rows`` dengan filter AI Status / Manual Status.

        Keduanya DITURUNKAN per hasil (tidak disimpan di baris), jadi tidak bisa
        difilter di SQL — harus dihitung dengan helper kanonik yang sama seperti
        yang dipakai tabelnya, lalu disaring di Python.

        Dipakai BERSAMA oleh menu Results, Manual Check, dan Pending
        Check. Sebelumnya logika ini hanya ada di cabang Results, sehingga di kedua
        menu antrean itu dropdown "Semua AI Status" / "Semua Manual Status" tetap
        tampil tetapi tidak berpengaruh apa pun — pengguna menyaring, jumlahnya
        tidak berubah, tanpa satu pun petunjuk kenapa.

        Hasil yang belum ``done`` dibuang saat filter aktif: AI Status baru ada
        setelah evaluasi selesai, jadi memasukkannya berarti menampilkan baris yang
        tidak mungkin cocok dengan nilai mana pun.
        """
        if not _has_status_filter:
            return rows
        rows = [r for r in rows if r.status == "done"]
        if not rows:
            return rows
        ids = [str(r.id) for r in rows]
        qcs = crud.qc_status_requests_for(db, ids)
        # Vonis dihitung helper KANONIK berkelompok — bahan yang sama persis dengan
        # kolom AI Status di tabelnya, ``doc_status`` termasuk. Menyusun bahannya
        # sendiri di sini pernah membuat filter melewatkan status dokumen, sehingga
        # tiket PENDING ikut muncul saat menyaring "Not Qualified".
        ai_map = ai_status_map(db, rows)
        out = []
        for r in rows:
            rid = str(r.id)
            ai = ai_map.get(rid)
            if ai_filter and ai != ai_filter:
                continue
            if manual_filter and effective_manual_status(qcs.get(rid), ai) != manual_filter:
                continue
            out.append(r)
        return out

    if banding_pending:
        # "Manual Check" (banding) menu, filtered per role:
        #  - QC: every ticket that HAS a banding (any status) so the QC can track the
        #    review outcome of the bandings it filed;
        #  - Team Leader QC: tickets with a banding still pending its check;
        #  - SPQ Head: tickets with a banding escalated to it.
        # Derived from appeals (not stored on the row), so load the scoped set and
        # filter in Python.
        all_results, _ = crud.list_results(
            db, campaign=campaign, campaigns=role_campaigns, ticket_id=ticket_id, page=1, limit=1_000_000,
            customer_ids=scoped_cids, date_start=d_start, date_end=d_end, **_iso,
        )
        appeal_all = crud.error_code_appeals_for_results(db, [str(r.id) for r in all_results])
        is_spq = has_perm(db, current_user, MANUAL_STATUS_REVIEW_SPQ) or has_perm(db, current_user, ERROR_CODE_REVIEW_SPQ)
        # Pengaju banding melihat SEMUA banding-nya (apa pun statusnya) agar bisa
        # memantau hasil review; reviewer hanya melihat yang menunggu gilirannya.
        is_qc = has_perm(db, current_user, ERROR_CODE_APPEAL)
        def _needs_review(rid):
            s = _appeal_summary(appeal_all.get(rid) or [])
            if not s:
                return False
            if is_qc:
                return s["total"] > 0
            return s["spq_pending"] > 0 if is_spq else s["tl_pending"] > 0
        matched = _apply_status_filters([r for r in all_results if _needs_review(str(r.id))])
        total = len(matched)
        results = matched[(page - 1) * limit : page * limit]
    elif manual_status_pending:
        # "Pending Check" menu (TL QC / SPQ Head): tickets whose Manual Status COLUMN
        # is "pending" — a QC status request still awaiting the hierarchy OR the
        # missing-documents default. Uses the SAME _manual_status helper the column
        # renders, so the queue matches the column value exactly (not stored on the
        # row, so load the scoped set and filter in Python).
        all_results, _ = crud.list_results(
            db, campaign=campaign, campaigns=role_campaigns, ticket_id=ticket_id, page=1, limit=1_000_000,
            customer_ids=scoped_cids, date_start=d_start, date_end=d_end, **_iso,
        )
        all_ids = [str(r.id) for r in all_results]
        qc_all = crud.qc_status_requests_for(db, all_ids)
        mdocs_all = _missing_docs_map(db, all_results)
        matched = _apply_status_filters([
            r for r in all_results
            if manual_review_state(qc_all.get(str(r.id)), mdocs_all.get(str(r.id), False)) == "menunggu"
        ])
        total = len(matched)
        results = matched[(page - 1) * limit : page * limit]
    elif _has_status_filter:
        # Menu Results dengan filter AI/Manual Status aktif: muat seluruh hasil `done`
        # yang cocok, lalu saring dengan helper yang SAMA dipakai kedua menu antrean di
        # atas — supaya "Qualified" berarti hal yang sama persis di ketiga menu.
        all_results, _ = crud.list_results(
            db, status="done", campaign=campaign, campaigns=role_campaigns, ticket_id=ticket_id, page=1,
            limit=1_000_000, customer_ids=scoped_cids, date_start=d_start, date_end=d_end, **_iso,
        )
        matched = _apply_status_filters(all_results)
        total = len(matched)
        results = matched[(page - 1) * limit : page * limit]
    else:
        results, total = crud.list_results(
            db, status=status, campaign=campaign, campaigns=role_campaigns, ticket_id=ticket_id, page=page,
            limit=limit, customer_ids=scoped_cids, date_start=d_start, date_end=d_end, **_iso,
        )
    return results, total


@router.get("/list_results", response_model=ResultListResponse)
def list_results(
    status: Optional[str] = Query(None),
    campaign: Optional[str] = Query(None),
    ticket_id: Optional[str] = Query(None),
    ai_status: Optional[str] = Query(None, description="Filter AI Status: PASS (Qualified) | FAIL (Not Qualified) | PENDING"),
    manual_status: Optional[str] = Query(None, description="Filter Manual Status (vonis human; default mengikuti AI Status): PASS | FAIL | PENDING"),
    am_nip: Optional[str] = Query(None, description="Hierarchy filter: Area Manager NIP"),
    tl_nip: Optional[str] = Query(None, description="Hierarchy filter: Team Leader NIP"),
    agent_nip: Optional[str] = Query(None, description="Hierarchy filter: Sales Agent (TLO) NIP"),
    qc_username: Optional[str] = Query(None, description="Team Leader QC filter: tickets assigned to this QC (NIP)"),
    qc_support_username: Optional[str] = Query(None, description="Team Leader QC filter: tickets uploaded by this QC Support (NIP)"),
    date_start: Optional[str] = Query(None, description="Transcript-date lower bound (YYYY-MM-DD, WIB)"),
    date_end: Optional[str] = Query(None, description="Transcript-date upper bound (YYYY-MM-DD, WIB)"),
    banding_pending: bool = Query(False, description="Keep only tickets with an appeal awaiting the caller's review tier"),
    manual_status_pending: bool = Query(False, description="Pending Check menu: keep only tickets with a Manual Status request awaiting the caller's review tier (TL QC / SPQ Head)"),
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    results, total = _resolve_filtered_results(
        db, current_user,
        status=status, campaign=campaign, ticket_id=ticket_id, ai_status=ai_status,
        manual_status=manual_status, am_nip=am_nip, tl_nip=tl_nip, agent_nip=agent_nip,
        qc_username=qc_username, qc_support_username=qc_support_username,
        date_start=date_start, date_end=date_end, banding_pending=banding_pending,
        manual_status_pending=manual_status_pending, page=page, limit=limit,
    )
    # Customer Name + Nomor Kartu columns are surfaced to the "simple viewer" roles
    # (individual agent + team leader + area manager + telesales head), so only look
    # them up for those roles.
    is_simple_viewer = getattr(current_user, "role", None) in ("sales_agent", "team_leader", "area_manager", "telesales_head")
    # [FIX] Result id halaman ini — dipakai beberapa query batched di bawah.
    # Definisi ini hilang saat enrichment dipindah dari _build_items() ke sini,
    # menyebabkan NameError pada doctimes/qc_map.
    rid_list = [str(r.id) for r in results]
    # Which results on this page already have uploaded documents, plus the latest
    # upload time per result (single grouped query, no N+1).
    doctimes = crud.document_upload_times(db, rid_list)
    docset = set(doctimes.keys())
    # Which document TYPES are already uploaded per result — a similarity band asks
    # for one specific document, so "has some document" is not enough.
    doctypes_map = crud.document_types_by_result(db, [str(r.id) for r in results])
    # QC-proposed AI-status changes for this page (single grouped query, no N+1).
    qc_map = crud.qc_status_requests_for(db, rid_list)
    # Error Code appeals for this page (single grouped query, no N+1). Approved
    # appeals lift the score by flipping the linked scorecard item to SESUAI.
    appeal_map = crud.error_code_appeals_for_results(db, [str(r.id) for r in results])
    # TMS data-change flags per result (keyed by cid = customer id from the source
    # filename prefix). Gates the Upload Document button. Single batched query.
    cids = [_customer_id_from_files(r.source_files) for r in results]
    change_map = crud.get_tms_cashline_change_flags(db, [c for c in cids if c])
    # TMS submit_time per cid — basis for the Pending Check H+2 SLA timer (counts
    # from the disbursement submission time, not the transcript's generated_at).
    submit_map = crud.tms_submit_time_map(db, [c for c in cids if c])
    # QC ticket -> (assignee, assigned_at). Team Leader QC / SPQ Head see who a
    # ticket is assigned to and when ("Assign Date").
    assignment_map = crud.assignment_map_for_tickets(db, [c for c in cids if c])
    # Tiket di halaman ini yang reprosesnya masih mengantre/berjalan (satu query).
    # Menahan tombol Reprocess sesudah refresh atau pindah menu: layar hanya
    # mengingat job yang ia mulai sendiri di sesi itu, jadi tanpa penanda dari
    # server tombolnya kembali enable padahal server akan menolaknya dengan 409.
    reprocessing_cids = crud.active_reprocess_ticket_ids(db, [c for c in cids if c])
    # Per-ticket manual checks by QC for this page (batched, no N+1).
    manual_check_map = crud.qc_manual_checks_for_results(db, [str(r.id) for r in results])
    # Riwayat Manual Status per tiket (append-only) — hanya jumlahnya yang dikirim ke
    # tabel; detailnya diambil per tiket lewat GET /qc_status_events/{result_id}.
    ms_events_map = crud.qc_status_events_for_results(db, [str(r.id) for r in results])
    # Which results are "missing required documents" (TMS data changed or limit >= 50jt
    # but nothing uploaded) — drives the AI-status default + Manual Status "pending".
    mdocs_page = _missing_docs_map(db, results)
    gaps_page = data_gap_map(db, results)
    # Keterangan LENGKAP di kolom AI Status hanya untuk pemegang capability-nya
    # (divisi QC + Admin). Empat role sisi sales hanya menerima keterangan
    # "Menunggu dokumen ..." — satu-satunya alasan yang bisa mereka tindaklanjuti
    # sendiri. Disaring di SINI, bukan di Vue: alasan yang disembunyikan tidak ikut
    # terkirim ke browser sama sekali.
    show_full_reason = has_perm(db, current_user, RESULTS_STATUS_REASON_FULL)
    doc_status_page = document_status_map(db, results)
    _now_status = datetime.now()  # basis tenggat H+2 untuk status PENDING
    items = []
    for r in results:
        ai_score = None
        ai_status = None
        passing_grade = None  # dynamic: taken from the LLM evaluation below
        maximum_score = None
        audio_duration = None
        audio_durations = None
        excluded_calls = None
        campaign_interest = None
        critical_compliance_check = None
        score_bomb = None
        non_tolerable_items = None
        scorecard_issues = []
        evaluation = None  # per-iteration; the RETURN rule below reads it after the QC override
        raw_result_json = None  # pre-appeal JSON — basis for the similarity-band doc triggers
        if r.status == "done":
            data = crud.get_result_data(db, str(r.id))
            if data is not None:
                raw_result_json = data.result_json
                if isinstance(data.result_json, dict):
                    audio_duration = data.result_json.get("audio_duration")
                    audio_durations = data.result_json.get("audio_durations")
                    excluded_calls = data.result_json.get("excluded_calls")
                evaluation = _evaluation_dict(data.result_json)
                if evaluation is not None:
                    # Terapkan banding yang di-approve (baris error code dihapus &
                    # item scorecard -> SESUAI) sebelum menghitung skor/status.
                    _row_appeals = appeal_map.get(str(r.id), [])
                    _approved = approved_appeals_only(_row_appeals)
                    _flip = [a for a in appeals_that_flip(_approved) if _appeal_kind(a) != "add"]
                    # Zona abu-abu ditegakkan di kode, sebelum banding & skor dihitung.
                    evaluation = normalize_static_verification(evaluation)
                    # Baris verifikasi dinamis tanpa dua sisi pembanding (Ascend/transkrip kosong)
                    # turun ke SKIPPED_NULL, bukan MISMATCH — tidak menerbitkan B17.
                    evaluation = normalize_dynamic_verification(evaluation)
                    # Lalu status dokumennya: baris zona abu-abu jadi PENDING selama
                    # tenggat H+2 berjalan, dan MISMATCH bila terlewat tanpa unggah.
                    _ds = doc_status_page.get(str(r.id))
                    if _ds is not None:
                        evaluation = apply_static_document_status(evaluation, _ds[0], _ds[1])
                        # Sisi cashline menyusul: baris yang menunggu cover buku tabungan menjadi
                        # PENDING selama tenggat H+2, agar jalur dokumennya tidak terpotong.
                        evaluation = apply_cashline_document_status(evaluation, _ds[0], _ds[1])
                    evaluation = apply_approved_appeals(evaluation, _flip)
                    evaluation = apply_approved_card_holder_appeals(evaluation, _flip)
                    evaluation = apply_approved_cashline_appeals(evaluation, _flip)
                    evaluation = apply_approved_critical_compliance_appeals(evaluation, _flip)
                    # 'add' bandings lower the score (attach a new error).
                    evaluation = apply_added_score_appeals(evaluation, added_appeals_only(_row_appeals))
                    # Non-tolerable (tolerable=NO) scorecard items still BELUM_SESUAI —
                    # surfaced as a Results column (and what forces AI Status = RETURN).
                    non_tolerable_items = _non_tolerable_reasons(evaluation)
                    scorecard_issues = _scorecard_issues(evaluation)

                    # AI Score dihitung deterministik dari scorecard (sumber tunggal,
                    # konsisten dengan XLSX export): phase2 = skor maksimal dikurangi
                    # bobot item BELUM_SESUAI; phase3 = phase2 + verifikasi + kritis.
                    # Satu sumber rumus (termasuk iris 10% non-tolerable) —
                    # lihat compliance.scoring.phase3_score.
                    score_total = _phase3_score(evaluation)
                    if score_total is not None:
                        ai_score = int(score_total) if score_total == int(score_total) else score_total

                    eval_pg = _numeric_or_none(evaluation.get("passing_grade"))
                    if eval_pg is not None:
                        passing_grade = eval_pg
                    maximum_score = _numeric_or_none(evaluation.get("maximum_score"))

                    # AI Status mengikuti skor deterministik di atas vs passing grade;
                    # fallback ke ai_status dari LLM bila skor/passing tidak tersedia.
                    if ai_score is not None and passing_grade is not None:
                        ai_status = "PASS" if ai_score >= passing_grade else "FAIL"
                    else:
                        status_val = evaluation.get("ai_status")
                        if isinstance(status_val, str):
                            ai_status = status_val.strip().upper()

                    ci = evaluation.get("campaign_interest")
                    if isinstance(ci, list):
                        campaign_interest = ci
                    # One sentence per failed critical item, computed server-side so the
                    # Results column and the detail view read identically (and the static
                    # verification items say WHY they failed instead of being negated).
                    evaluation = annotate_critical_compliance_reasons(evaluation)
                    # Kolom SCOREBOMB: SELURUH item yang mengebom skor — kritis (25%)
                    # DAN non-tolerable lain (10%). Sampai 31 Agustus 2026 kolom itu
                    # hanya memuat yang kritis, sehingga iris 10% memotong skor tanpa
                    # pernah terlihat.
                    score_bomb = _score_bomb_items(evaluation)
                    ccc = evaluation.get("critical_compliance_check")
                    if isinstance(ccc, dict):
                        # Keep only what the Results column needs: overall status +
                        # the checked items (item_code / requirement / status / reason).
                        items_raw = ccc.get("checked_items")
                        critical_compliance_check = {
                            "status": ccc.get("status"),
                            "checked_items": items_raw if isinstance(items_raw, list) else [],
                        }
        cid = _customer_id_from_files(r.source_files)
        customer_name = None
        account_number = None
        credit_limit = None
        if cid:
            # Fetch the customer's data once: simple viewers need name/card too, the
            # rest need only the credit limit (for the NPWP-upload trigger).
            if is_simple_viewer:
                info, _ = get_customer_info(cid, db)
                customer_name = info.get("customer_name")
                account_number = info.get("account_number")
                credit_limit = info.get("credit_limit")
            else:
                credit_limit = get_credit_limit(cid, db)
        # Upload Document gate: dokumen apa yang diminta tiket ini dan KARENA APA —
        # perubahan data TMS (Alamat Kantor/Rumah, NPWP, NIK), band similarity card
        # holder (Fase #5), atau limit pencairan >= Rp 50 juta. Band-nya dibaca dari
        # JSON PRA-banding supaya sepakat dengan _missing_docs_map, yang menentukan
        # status PENDING.
        #
        # PRA-BANDING, TAPI SUDAH DINORMALKAN (1 September 2026). Sebelumnya
        # ``raw_result_json`` dioper apa adanya, padahal ``_missing_docs_map`` sudah
        # memakai ``_normalized_json`` sejak 21 Agustus 2026 — jadi justru TIDAK
        # sepakat, dan kolom Document memakai angka mentah LLM sementara status/skor
        # memakai hasil hitung ulang Python. Dua arah salahnya sama-sama nyata:
        # 020338gGlU meminta KK atas nama ibu kandung "67%" yang sebenarnya 90%
        # (MATCH bersih, tiket tetap Qualified), sedangkan 020532VF8Y berstatus
        # PENDING karena dokumennya kurang tetapi kolomnya kosong sehingga slot
        # unggahnya tidak pernah terbuka. Normalisasi TIDAK sama dengan banding:
        # ia hanya menegakkan band di Python, jadi sifat "pra-banding" tetap utuh.
        flags = change_map.get(cid or "", {})
        doc_requirements = _document_requirements(
            flags,
            _normalized_json(raw_result_json),
            card_holder_bands_apply(r.uploaded_at),
            npwp_required_by_limit(credit_limit),
        )
        # Dua daftar lama tetap dikirim (modal upload & payload lama memakainya),
        # sekarang diturunkan dari satu sumber di atas alih-alih dihitung ulang.
        document_triggers = list(dict.fromkeys(d["reason"] for d in doc_requirements))
        document_upload_types = list(dict.fromkeys(d["doc_type"] for d in doc_requirements))
        document_missing_types = [
            t for t in document_upload_types
            if t not in doctypes_map.get(str(r.id), set())
        ]
        # "Kekurangan dokumen": dokumen wajib (perubahan data TMS ATAU limit >= 50jt)
        # tapi belum ada upload. Basis sama dengan chart/KPI (_missing_docs_map).
        missing_docs = mdocs_page.get(str(r.id), False)

        # --- AI Status (aturan 7 Agustus 2026) ---
        # Skor di atas TETAP dihitung dari evaluasi + banding error code apa adanya —
        # kolom AI Score, tabel Error Code, dan Ringkasan Kategori memang harus jujur.
        # Yang dikunci hanyalah STATUS-nya.
        qc_req = qc_map.get(str(r.id))
        # 1) non-tolerable veto — tanpa syarat status sebelumnya (28 Agustus 2026).
        blocking = evaluation is not None and _has_blocking_intolerable_item(evaluation)
        if blocking:
            ai_status = "FAIL"
        # 2) missing-documents: dalam tenggat H+2 -> PENDING; lewat tenggat -> FAIL.
        # ``and not blocking``: tiket yang sudah kena pelanggaran non-tolerable tidak
        # singgah di PENDING — dokumen susulan tidak akan menggugurkan pelanggaran
        # yang sudah terbukti dari transkrip. Urutan ini WAJIB sama dengan helper
        # kanonik ``_result_ai_status``; kalau berbeda, kolom di layar bertentangan
        # dengan angka di Statistics dan dengan filter AI Status.
        if missing_docs and not blocking:
            ai_status = "FAIL" if _doc_sla_expired(submit_map.get(cid), _now_status) else "PENDING"
        data_gap = gaps_page.get(str(r.id)) or ()
        # 2b) KONSISTENSI VERIFIKASI STATIK: gugur di TAHAP 1 (penyebutan nasabah
        # tidak konsisten antar pengulangan). Diperiksa SESUDAH aturan dokumen karena
        # harus menimpa PENDING — tiket seperti ini tidak menunggu dokumen.
        fraud_fields = static_consistency_failures(evaluation) if evaluation is not None else []
        if fraud_fields:
            ai_status = "FAIL"
        # 2c) BADWORD (13 Agustus 2026): agent mengucapkan kalimat bersentimen negatif
        # kepada nasabah — perilaku yang tidak dapat ditoleransi dan sumber komplain
        # CCBM. Sama seperti aturan 2b: menimpa PENDING (tiket seperti ini tidak
        # menunggu dokumen) dan tidak peduli berapa skornya.
        badwords = badword_rows(evaluation) if evaluation is not None else []
        if badwords:
            ai_status = "FAIL"
        # 2d) KEKURANGAN DATA ACUAN (28 Agustus 2026): transkrip / TMS / agent /
        # Ascend yang kosong membuat tiket PENDING, menimpa SEMUA aturan di atas
        # (veto non-tolerable, tenggat dokumen, konsistensi statik, badword) — skor
        # maupun temuan yang memicu keempatnya lahir dari evaluasi yang acuannya tidak
        # lengkap. Urutannya sama persis dengan helper kanonik ``_result_ai_status``;
        # kalau berbeda, kolom di layar akan bertentangan dengan angka di Statistics.
        if data_gap:
            ai_status = "PENDING"
        # 3) Vonis human yang SUDAH DISETUJUI mengunci AI Status — menimpa ketiga aturan
        # di atas. Sengaja memanggil helper kanonik yang sama dengan yang dipakai Stats,
        # filter, dan snapshot: baris ini dulu menghitung sendiri tanpa melihat qc_req,
        # sehingga kolomnya bisa berbunyi Qualified sementara filter "AI Status =
        # Not Qualified" justru memasukkannya.
        _override = manual_status_of(qc_req)
        if _override:
            ai_status = _override
        # "Manual Status" column (all roles): vonis human bila ada, selain itu mengikuti
        # AI Status. `manual_status_by_human` membedakan keduanya (tombol Set vs Ubah).
        manual_status = _manual_status(qc_req, ai_status)
        manual_by_human = manual_status_of(qc_req) is not None
        manual_review = manual_review_state(qc_req, missing_docs)
        # Kenapa AI Status-nya PENDING. Satu-satunya penyebab PENDING yang dihasilkan
        # sistem adalah dokumen wajib yang belum diunggah (tenggat H+2 dari
        # tms_cashline.submit_time); PENDING yang datang langsung dari LLM tidak punya
        # keterangan, jadi dibiarkan kosong ketimbang mengarang alasan.
        # Badword ditulis sebagai alasan tiket, sejajar pending_reason, supaya QC
        # melihat SEBABNYA di kolom AI Status tanpa membuka baris.
        # Kegagalan konsistensi verifikasi statik SENGAJA tidak diberi catatan di sini
        # (kebijakan 21 Agustus 2026): aturannya tetap menggugurkan tiket, tetapi
        # sebabnya dibaca dari kolom Critical Failure dan baris verifikasinya, bukan
        # dari label khusus di kolom AI Status.
        fail_notes = [] if show_full_reason else None
        # Alasan Not Qualified yang TIDAK terbaca dari skor. Sengaja TANPA daftar
        # item_code (permintaan 28 Agustus 2026): item mana yang gagal sudah terbaca
        # di tabel Error Code saat barisnya dibuka, jadi menuliskannya lagi di kolom
        # status hanya memanjangkan sel tanpa menambah informasi.
        if fail_notes is not None:
            if blocking and ai_status == "FAIL":
                fail_notes.append("Error non-tolerable")
            if badwords and ai_status == "FAIL":
                fail_notes.append(badword_fail_reason(evaluation))
        fail_reason = " · ".join(n for n in (fail_notes or []) if n) or None
        pending_reasons = []
        if ai_status == "PENDING" and data_gap and show_full_reason:
            # Ditulis PALING DEPAN: ini penyebab yang tidak bisa diselesaikan agent
            # maupun QC — datanya harus dilengkapi lebih dulu. SEMUA kekurangan
            # disebut, bukan hanya yang pertama ditemukan.
            pending_reasons.append(data_gap_reasons(data_gap))
        if ai_status == "PENDING" and missing_docs:
            _need = [_doc_label(t) for t in document_missing_types] or ["pendukung"]
            # Tenggatnya hanya disebut bila aturannya memang aktif — menuliskan
            # "SLA H+2" saat aturan itu dimatikan justru menyesatkan.
            _sla = " (SLA H+2)" if doc_sla_enabled() else ""
            pending_reasons.append(f"Menunggu dokumen {', '.join(_need)}{_sla}")
        # Sebuah tiket bisa PENDING karena DUA hal sekaligus; keduanya ditulis apa
        # adanya ketimbang memilih salah satu dan menyembunyikan sisanya.
        pending_reason = " · ".join(pending_reasons) or None
        items.append(
            ResultListItem(
                result_id=str(r.id),
                id=cid,
                customer_name=customer_name,
                account_number=account_number,
                credit_limit=(credit_limit if is_simple_viewer else None),
                campaign=r.campaign,
                source_files=r.source_files,
                num_calls=r.num_calls,
                audio_duration=audio_duration,
                audio_durations=audio_durations,
                excluded_calls=excluded_calls,
                campaign_interest=campaign_interest,
                critical_compliance_check=critical_compliance_check,
                score_bomb_items=score_bomb,
                non_tolerable_items=non_tolerable_items,
                scorecard_issues=scorecard_issues,
                status=r.status,
                ai_score=ai_score,
                passing_grade=passing_grade,
                maximum_score=maximum_score,
                ai_status=ai_status,
                uploaded_at=r.uploaded_at,
                uploaded_by_username=r.uploaded_by_username,
                uploaded_by_role=r.uploaded_by_role,
                generated_at=r.generated_at,
                submit_time=(submit_map.get(cid) or None),
                completed_at=r.completed_at,
                processing_sec=r.processing_sec,
                document_triggers=document_triggers,
                document_upload_types=document_upload_types,
                document_missing_types=document_missing_types,
                document_missing_labels=[_doc_label(t) for t in document_missing_types],
                document_requirements=doc_requirements,
                reprocess_active=(cid in reprocessing_cids) if cid else False,
                has_documents=str(r.id) in docset,
                document_uploaded_at=doctimes.get(str(r.id)),
                qc_request=qc_req,
                manual_status=manual_status,
                manual_review_state=manual_review,
                manual_status_by_human=manual_by_human,
                missing_documents=missing_docs,
                pending_reason=pending_reason,
                fail_reason=fail_reason,
                manual_status_history_count=len(ms_events_map.get(str(r.id), [])),
                appeal_summary=_appeal_summary(appeal_map.get(str(r.id), [])),
                assigned_qc=(assignment_map.get(cid or "") or (None, None))[0],
                assigned_at=(assignment_map.get(cid or "") or (None, None))[1],
                qc_checked_at=getattr(manual_check_map.get(str(r.id)), "created_at", None),
                qc_checked_by=getattr(manual_check_map.get(str(r.id)), "checked_by_username", None),
            )
        )
    return ResultListResponse(items=items, total=total, page=page, limit=limit)


@router.get("/stats", response_model=StatsResponse)
def get_stats(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    data = crud.get_stats(
        db,
        _scoped_customer_ids(db, current_user),
        effective_campaigns_for(db, current_user),
    )
    return StatsResponse(**data)


@router.get("/stats/daily", response_model=DailyStatsResponse)
def get_daily_stats(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    days = crud.get_daily_stats(db, _scoped_customer_ids(db, current_user))
    return DailyStatsResponse(days=days)


# --- Revamped Statistics dashboard (daily-cached snapshot) ------------------
# All three endpoints read the same once-per-day snapshot (overview, per-agent,
# campaign month-to-month, org hierarchy); see crud.get_or_build_stats_snapshot.

def _snapshot_for(db: Session, current_user) -> dict:
    """Statistics snapshot sesuai scope pemakainya.

    Area Manager mendapat snapshot yang dihitung ulang dari tiket area-nya saja,
    sehingga KPI, donut, Performa Sales, Performa Campaign dan hierarki semuanya
    konsisten satu sama lain. Role lain memakai snapshot global yang di-cache.

    Catatan: versi scoped TIDAK di-cache — dihitung per request. Bebannya
    terbatas karena query-nya sudah difilter di sisi DB ke tiket area tersebut.
    """
    scope = data_scope_for(db, current_user)
    if scope in (SCOPE_QC_SUPPORT_OWN, SCOPE_QC_ASSIGNED):
        # Kedua cakupan QC perorangan berdiri sendiri: seluruh halaman Statistics
        # dihitung dari himpunan tiketnya saja — QC Support dari tiket complaint yang
        # ia unggah, QC dari tiket yang di-ASSIGN kepadanya. Tanpa ini keduanya
        # menerima snapshot global, sehingga Statistics mengaku 102 tiket sementara
        # menu Results-nya hanya memuat 6 (kasus nyata user ``bella``, 7 Agustus 2026)
        # — angka yang tidak bisa ditelusuri ke satu barisnya pun.
        # ``roster_uids`` sengaja himpunan KOSONG supaya tidak ada seorang pun di-seed
        # dari roster: tabel Performa Sales-nya hanya boleh memuat agent yang benar-benar
        # muncul di tiketnya (kalau tidak, 262 baris agent seluruh organisasi ikut tampil).
        return compute_stats_snapshot(
            db, _scoped_customer_ids(db, current_user) or [], roster_uids=set()
        )
    if not is_sales_scope(scope):
        # Cakupan non-sales yang DIBATASI CAMPAIGN ikut dihitung ulang: snapshot
        # global yang di-cache memuat seluruh organisasi, jadi memakainya membuat
        # role ber-``data_scope: all`` + tag campaign melihat KPI, Performa Campaign
        # dan hierarki milik campaign yang bukan cakupannya (12 Agustus 2026 — role
        # bertag CECC masih membaca 98 tiket cashline). ``roster_uids`` kosong dengan
        # alasan yang sama dengan cabang QC di atas.
        cids = _scoped_customer_ids(db, current_user)
        if cids is None:
            return crud.get_or_build_stats_snapshot(db)
        return compute_stats_snapshot(db, cids, roster_uids=set())

    # KETIGA cakupan sales dihitung ulang dari tiketnya sendiri. Sebelumnya hanya
    # ``sales_am`` yang di-scope, sehingga Team Leader & Sales Agent menerima snapshot
    # GLOBAL dari endpoint ini — termasuk daftar ``agents`` berisi seluruh organisasi.
    # Tidak terlihat di UI (keduanya memakai tampilan my_overview), tetapi tetap
    # terkirim di respons, dan jumlahnya membengkak setelah Performa Sales di-seed
    # dari roster.
    username = getattr(current_user, "username", "") or ""
    campaigns = effective_campaigns_for(db, current_user)
    if scope == SCOPE_SALES_AM:
        roster_uids = agent_ids_for_am(db, username, campaigns)
    elif scope == SCOPE_SALES_TL:
        roster_uids = agent_ids_for_tl(db, username, campaigns)
    else:
        roster_uids = agent_ids_for_agent(db, username, campaigns)

    # Seeding roster dibatasi ke agent dalam cakupan ini saja — tanpa itu, Performa
    # Sales & pohon hierarki akan memuat orang di luar cakupannya (dengan angka nol).
    return compute_stats_snapshot(
        db,
        _scoped_customer_ids(db, current_user) or [],
        roster_uids=roster_uids,
    )


@router.get("/stats/overview")
def stats_overview(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    """KPI cards + status donut + sales-performance table."""
    snap = _snapshot_for(db, current_user)
    return {
        "overview": snap["overview"],
        "agents": snap["agents"],
        # Per-campaign breakdowns for the Overview campaign filter (older cached
        # snapshots may predate these keys, hence the defaults).
        "campaigns": snap.get("campaigns", []),
        "overview_by_campaign": snap.get("overview_by_campaign", {}),
        "agents_by_campaign": snap.get("agents_by_campaign", {}),
    }


@router.get("/stats/campaigns_monthly")
def stats_campaigns_monthly(
    db: Session = Depends(get_db), current_user=Depends(get_current_user)
):
    """Per (campaign, month) performance report: submissions, Risk Base breakdown
    (High/Medium/Low/System/New) and error rate. Area Manager melihat area-nya saja."""
    return _snapshot_for(db, current_user)["campaign_monthly"]


#: Field Risk Base pada tiap simpul pohon hierarki. Digerbangi ``STATS_RISK_BASE``
#: (H/M/L) dan ``STATS_RISK_SYSTEM_NEW`` (System/New) — capability yang sama dengan
#: yang menyembunyikan kolomnya di layar.
#:
#: ``total_risk`` SENGAJA TIDAK ADA DI SINI sejak 28 Agustus 2026 sore. Sebelumnya ia
#: ikut dibuang bersama H/M/L, sehingga sisi sales menerima ``error_rate`` tanpa
#: pembilangnya. Begitu kolom "Errors" diganti nama menjadi "Total Failure" atas
#: permintaan bisnis, kolom itu terpaksa diisi ``errors`` (jumlah TIKET gagal) —
#: angka yang berbeda dari pembilang rasionya, sehingga
#: ``Total Failure ÷ Submissions`` tidak sama dengan ``Failure Rate`` yang tertera
#: (mis. 26 / 98 = 26,5% padahal yang tampil 68,4%).
#:
#: Yang ditahan gate ini memang RINCIAN risk base (pecahan per severity), bukan
#: totalnya: totalnya sudah lama terbaca sisi sales dalam bentuk persentase. Jadi
#: mengirim ``total_risk`` tidak membuka informasi baru, hanya membuat kolomnya
#: menamai angka yang benar.
_HIER_RISK_FIELDS = ("risk_high", "risk_medium", "risk_low")
_HIER_RISK_EXTRA_FIELDS = ("risk_system", "risk_new")


def _strip_hierarchy_risk(node, drop: tuple):
    """Buang field Risk Base dari sebuah simpul pohon (rekursif ke bawah).

    Sampai 28 Agustus 2026 angka Risk Base SELALU ikut terkirim, dan yang menyembunyikannya
    hanya ``v-if`` di Vue — artinya sisi sales tetap bisa membacanya lewat DevTools
    meski capability-nya tidak mereka punya. Capability yang dipasang untuk menahan
    informasi harus menahannya di server, bukan sekadar tidak menggambarnya.

    ``error_rate`` dan ``total_risk`` sengaja DIPERTAHANKAN: rasionya memang
    ditampilkan ke sisi sales, dan sejak 28 Agustus 2026 pembilangnya ikut — lihat
    catatan pada ``_HIER_RISK_FIELDS``. Yang ditahan tinggal pecahan per severity
    (High/Medium/Low, plus System/New untuk gate satunya).
    """
    if not isinstance(node, dict):
        return node
    for key in drop:
        node.pop(key, None)
    # ``all_telesales`` sebuah dict tunggal, sisanya list — keduanya harus ditelusuri.
    # Melewatkan ``all_telesales`` berarti membocorkan justru angka yang paling besar:
    # total Risk Base seluruh organisasi.
    for child_key in ("all_telesales", "area_managers", "team_leaders", "agents", "tickets"):
        child = node.get(child_key)
        if isinstance(child, dict):
            _strip_hierarchy_risk(child, drop)
        else:
            for item in child or ():
                _strip_hierarchy_risk(item, drop)
    return node


def _scoped_hierarchy(db, current_user, tree: dict) -> dict:
    """Pohon hierarki yang sudah dibersihkan dari field di luar hak pemanggil."""
    drop = ()
    if not has_perm(db, current_user, STATS_RISK_BASE):
        drop += _HIER_RISK_FIELDS + _HIER_RISK_EXTRA_FIELDS
    elif not has_perm(db, current_user, STATS_RISK_SYSTEM_NEW):
        drop += _HIER_RISK_EXTRA_FIELDS
    if not drop:
        return tree
    # deepcopy: payload-nya berasal dari snapshot yang di-cache di DB dan dipakai
    # bersama seluruh pemanggil — memodifikasinya di tempat akan merusak respons role lain.
    import copy

    return _strip_hierarchy_risk(copy.deepcopy(tree), drop)


@router.get("/stats/hierarchy")
def stats_hierarchy(
    campaign: Optional[str] = Query(None, description="Batasi pohon ke satu campaign"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Error-rate hierarchy: Area Manager -> Team Leader -> Agent, plus All Telesales.

    Telesales Head, SPQ Head/Admin dan TL QC melihat pohon penuh. Seorang
    ``area_manager`` hanya melihat node-nya sendiri (TL + agent di bawahnya),
    termasuk ``all_telesales`` yang ikut dihitung dari area itu saja.

    ``campaign`` (opsional) memakai pohon per campaign dari snapshot yang sama,
    dibangun fungsi yang sama dengan versi global — jadi angkanya konsisten. Nama
    campaign dicocokkan case-insensitive karena kunci snapshot berasal dari
    ``Result.campaign`` apa adanya, sedangkan dropdown dari daftar campaign aktif.
    """
    snap = _snapshot_for(db, current_user)
    if not campaign:
        return _scoped_hierarchy(db, current_user, snap["hierarchy"])
    by_camp = snap.get("hierarchy_by_campaign") or {}
    want = campaign.strip().casefold()
    for key, tree in by_camp.items():
        if (key or "").strip().casefold() == want:
            return _scoped_hierarchy(db, current_user, tree)
    # Campaign aktif yang belum punya tiket: pohon kosong, bukan pohon global.
    from qc_core.compliance.stats_aggregate import empty_hierarchy

    return empty_hierarchy()


@router.get("/stats/role_counts")
def stats_role_counts(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    """User counts per role, split into the Sales and QC divisions — for the
    "Jumlah Sales" / "Jumlah QC" tables under the Overview donut."""
    from sqlalchemy import func as _f
    from qc_core.db.models import User
    rows = dict(
        db.query(User.role, _f.count()).filter(User.is_active == True).group_by(User.role).all()  # noqa: E712
    )
    return {
        "sales": {r: rows.get(r, 0) for r in ("telesales_head", "area_manager", "team_leader", "sales_agent")},
        "qc": {r: rows.get(r, 0) for r in ("team_leader_qc", "qc", "qc_support")},
    }


@router.get("/stats/my_overview")
def stats_my_overview(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    """Scoped KPI overview for an Area Manager, Team Leader or Sales Agent
    (individual agent): totals, failed count, and error rate over only the tickets
    they may see. For a Team Leader (team) or Area Manager (area) the response also
    includes ``agents`` — the per-agent roster (for the Statistics table); each row
    carries its ``team_leader`` so the Area Manager roster can group by TL. For an
    Area Manager the response also carries ``hierarchy`` — a scoped Team Leader ->
    Agent tree for the "Hierarki Failure Rate" tab (no Area Manager level)."""
    from qc_core.compliance.stats_aggregate import (
        compute_scoped_hierarchy,
        compute_scoped_overview,
        compute_team_agents,
    )

    username = getattr(current_user, "username", "") or ""
    scope = data_scope_for(db, current_user)
    campaigns = effective_campaigns_for(db, current_user)
    cids = _scoped_customer_ids(db, current_user)
    resp = {"overview": compute_scoped_overview(db, cids or [])}
    if scope == SCOPE_SALES_AM:
        roster = compute_team_agents(db, agent_ids_for_am(db, username, campaigns))
        resp["agents"] = roster
        resp["hierarchy"] = compute_scoped_hierarchy(roster)
    elif scope == SCOPE_SALES_TL:
        resp["agents"] = compute_team_agents(db, agent_ids_for_tl(db, username, campaigns))
    return resp


@router.get("/stats/failure_reasons")
def stats_failure_reasons(
    campaign: Optional[str] = Query(None, description="Batasi ke satu campaign"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Kategori scorecard yang paling sering gagal + alasannya. Hanya untuk
    SPQ Head & Admin (tab "Failure Reason" di menu Stats)."""
    if not has_perm(db, current_user, STATS_FAILURE_REASON):
        raise HTTPException(status_code=403, detail="Role Anda tidak memiliki akses Failure Reason.")
    from qc_core.compliance.stats_aggregate import compute_failure_reasons
    return compute_failure_reasons(db, campaign, effective_campaigns_for(db, current_user))


@router.get("/stats/failure_reasons_hierarchy")
def stats_failure_reasons_hierarchy(
    campaign: Optional[str] = Query(None, description="Batasi ke satu campaign"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Failure Reason yang dipecah per hierarki sales (AM -> TL -> Agent): kategori
    scorecard terbesar MILIK tiap simpul. Sub-tab "Hierarki Based" pada tab Failure
    Reason; hak aksesnya sama dengan agregatnya."""
    if not has_perm(db, current_user, STATS_FAILURE_REASON):
        raise HTTPException(status_code=403, detail="Role Anda tidak memiliki akses Failure Reason.")
    from qc_core.compliance.stats_aggregate import compute_failure_reasons_hierarchy
    return compute_failure_reasons_hierarchy(db, campaign,
                                             effective_campaigns_for(db, current_user))


@router.get("/stats/ai_status_timeseries")
def stats_ai_status_timeseries(
    granularity: str = "monthly",
    start: Optional[str] = None,
    end: Optional[str] = None,
    campaign: Optional[str] = None,
    offset: int = 0,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Approve/Reject proportions over time for the 100% stacked column chart.

    Scoped exactly like the KPI/donut it feeds — memakai aturan yang sama dengan
    ``_snapshot_for``: Sales Agent / Team Leader / Area Manager melihat tiketnya
    sendiri, QC melihat tiket yang di-assign kepadanya, QC Support melihat tiket
    complaint-nya; TL QC / Telesales Head / SPQ Head / Admin melihat semuanya.
    ``granularity`` = daily|weekly|monthly|quarterly|semester|yearly;
    ``start``/``end`` = optional 'YYYY-MM-DD' WIB bounds; ``campaign`` optional filter;
    ``offset`` pages the default window by whole windows (0 = latest, <0 older)."""
    # ``_scoped_customer_ids`` sudah menjawab None untuk cakupan tanpa penyempitan,
    # jadi tidak perlu daftar role di sini — dan dengan begitu pembatasan CAMPAIGN
    # (yang juga berlaku pada ``data_scope: all``) ikut terbawa.
    customer_ids = _scoped_customer_ids(db, current_user)
    return compute_ai_status_timeseries(db, customer_ids, campaign, granularity, start, end, offset)


@router.post("/stats/refresh")
def stats_refresh(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    """Force-recompute today's Statistics snapshot (SPQ Head only)."""
    if not has_perm(db, current_user, STATS_FAILURE_REASON):
        raise HTTPException(status_code=403, detail="Role Anda tidak dapat me-refresh statistik")
    crud.get_or_build_stats_snapshot(db, force=True)
    return {"status": "ok"}


def _append_scorecard_rows(sheet, result_json, cust_id) -> None:
    """Append scorecard rows (one per ``scorecard_result`` item, plus one per
    ``verified_parameter`` for Verifikasi Dinamis) from a single result JSON.

    The ``id`` column is the result-level ``cust_id`` (not derived from ticket_id).

    Kolom ``category`` memakai nama TAMPILAN (``category_label``), sama dengan
    dashboard: "Verifikasi" ditulis "Verifikasi Statik". Perbandingan di bawah tetap
    memakai nilai MENTAH ``item["category"]`` — yang berubah hanya yang tercetak."""
    if not result_json:
        return
    evaluation = result_json.get("evaluation") or {}
    scorecard = evaluation.get("scorecard_result") or []
    for item in scorecard:
        evidence = item.get("evidence") or {}
        # LLM emits ticket_id inside evidence; fall back to an item-level field.
        ticket_id = evidence.get("ticket_id") or item.get("ticket_id") or ""
        ev_text = f'{evidence.get("timestamp", "")} {evidence.get("quote", "")}'.strip()
        sheet.append(
            [
                cust_id,
                category_label(item.get("category", "")),
                item.get("item_code", ""),
                item.get("requirement", ""),
                item.get("item_score", ""),
                ev_text,
                ticket_id,
            ]
        )
        # Verifikasi Dinamis: one row per verified_parameter (mirror dashboard).
        if item.get("category") == "Verifikasi Dinamis":
            detail = item.get("dynamic_verification_detail") or {}
            for param in detail.get("verified_parameters") or []:
                aq = param.get("agent_question") or {}
                ca = param.get("customer_answer") or {}
                p_ticket = aq.get("ticket_id") or ""
                agent = f'{aq.get("timestamp", "")} {aq.get("quote", "")}'.strip()
                customer = f'{ca.get("timestamp", "")} {ca.get("quote", "")}'.strip()
                p_evidence = f"Agent : {agent} | Customer Answer : {customer}"
                sheet.append(
                    [
                        cust_id,
                        category_label(item.get("category", "")),
                        item.get("item_code", ""),
                        param.get("param_name", ""),
                        "",  # score dikosongkan (SKOR)
                        p_evidence,
                        p_ticket,
                    ]
                )


# --- Error Code table -------------------------------------------------------
# All error-code logic lives in compliance/error_codes.py (single source of
# truth, shared with the dashboard via the /result API). This sheet renders the
# build_error_code_table() output, grouped by source ("sumber").


def _append_errorcard_rows(sheet, result_json, cust_id, appeals=None) -> None:
    """Append error-card rows mirroring the dashboard Error Code table:
    id | sumber | error_code | item_code | details_error | reason | evidence | ticket_id."""
    if not result_json:
        return
    evaluation = result_json.get("evaluation") or {}
    # Approved adds only — an export must not carry a banding that Team Leader QC /
    # SPQ Head has not accepted yet.
    # Sama dengan tampilan dashboard: B16 & B17-dinamis yang berulang jadi satu baris.
    table = inject_added_rows(
        merge_dynamic_verification_rows(build_error_code_table(evaluation)),
        added_appeals_only(appeals or []),
    )
    for row in table:
        sheet.append([
            cust_id,
            row["sumber"],
            row["error_code"],
            row["item_code"],
            row["details_error"],
            row["reason"],
            row["evidence"],
            row["ticket_id"],
        ])


def _scorecard_issues(evaluation: dict) -> list:
    """Item scorecard yang belum beres — isi kolom SCORECARD tata letak Demo.

    Dua keadaan, dan keduanya berbeda artinya:

    * ``BELUM_SESUAI`` — item benar-benar gagal; bobotnya SUDAH dipotong dari skor.
    * ``PENDING`` — item verifikasi statik (SC_CL_23_1 tanggal lahir / SC_CL_23_2 nama
      ibu kandung) yang baris verifikasinya masih di zona abu-abu: dokumen pendukungnya
      ditunggu dan tenggat H+2 belum lewat, jadi skornya BELUM dipotong. Begitu tenggat
      lewat tanpa unggah, ``apply_static_document_status`` menjadikan barisnya MISMATCH
      dan ``_propagate_verification_to_scorecard`` menurunkan item-nya ke BELUM_SESUAI —
      butir yang tadinya kuning berubah merah dengan sendirinya.

    Dibaca dari evaluasi yang SUDAH melewati status dokumen + banding, sama dengan
    sumber ``ai_score``, sehingga daftar ini tidak bisa berbeda dari skornya. Urutannya
    mengikuti urutan scorecard, bukan dikelompokkan per status.
    """
    pending_codes = set()
    for row in evaluation.get("card_holder_verification") or []:
        if not isinstance(row, dict) or row.get("match") != "PENDING":
            continue
        code = CARD_HOLDER_STATIC_SCORECARD.get(row.get("field"))
        if code:
            pending_codes.add(code)

    out = []
    for item in evaluation.get("scorecard_result") or []:
        if not isinstance(item, dict):
            continue
        code = str(item.get("item_code") or "").strip()
        if not code:
            continue
        status = str(item.get("status") or "").strip().upper()
        if status == "BELUM_SESUAI":
            out.append({"item_code": code, "status": "BELUM_SESUAI"})
        elif code in pending_codes:
            out.append({"item_code": code, "status": "PENDING"})
    return out


def _non_tolerable_reasons(evaluation: dict) -> list:
    """Negated reasons for non-tolerable (tolerable=NO) scorecard items still unmet
    (status=BELUM_SESUAI), e.g. "Agent tidak menjelaskan biaya" — TEXT ONLY, without
    the (SC_CL_x) code suffix. Reads the (appeal-adjusted) evaluation, so approved
    appeals that flip an item to SESUAI drop out of the list."""
    out = []
    for item in evaluation.get("scorecard_result") or []:
        tol = str((item or {}).get("tolerable") or "").strip().upper()
        st = str((item or {}).get("status") or "").strip().upper()
        if tol == "NO" and st == "BELUM_SESUAI":
            reason = not_fulfilled_reason(item, evaluation)  # "... (SC_CL_x)"
            code = (item or {}).get("item_code")
            if code:
                # Drop the " (SC_CL_x)" code tag wherever it appears (it sits mid-
                # string for the no-requirement fallback), leaving text only.
                reason = reason.replace(f" ({code})", "")
            out.append(reason)
    return out


def _has_blocking_intolerable_item(evaluation: dict) -> bool:
    """A scorecard item that is non-tolerable (tolerable=NO) yet still unmet
    (status=BELUM_SESUAI) forces an AI status of RETURN (FAIL), regardless of the
    numeric score. Reads the (appeal-adjusted) evaluation, so approved appeals
    that flip an item to SESUAI correctly stop triggering this."""
    return bool(_non_tolerable_reasons(evaluation))


# Label baris "Hasil" di sheet ringkasan, per AI Status. Sejajar dengan
# ``aiStatusLabel`` di dashboard (QUALIFIED / NOT QUALIFIED / PENDING), tetapi memakai
# istilah Indonesia yang sudah dipakai sheet ini sejak awal.
_HASIL_LABELS = {"PASS": "LULUS", "FAIL": "TIDAK LULUS", "PENDING": "PENDING"}


def _append_ringkasan_rows(sheet, result_json, ai_status=None, status_note=None) -> None:
    """Append the AI score-summary as a running-balance table (mirrors the
    dashboard "Ringkasan Penilaian AI"): Keterangan | Perubahan | Hasil.

    ``ai_status`` adalah vonis KANONIK tiket ini — hasil
    ``compliance.stats_aggregate.ai_status_for_result``, sumber yang sama dengan
    kolom AI Status di daftar Results. Sengaja dioper dari pemanggil alih-alih
    dihitung ulang di sini: sampai 14 Agustus 2026 sheet ini menghitung sendiri dari
    dict evaluasi saja, sehingga dua aturan yang butuh query DB tidak pernah
    terlihat olehnya — tenggat unggah dokumen H+2 (``documents`` +
    ``tms_cashline.submit_time``) dan Manual Status yang sudah di-approve
    (``qc_status_requests``). Akibatnya 21 dari 98 tiket ter-ekspor "LULUS" padahal
    daftar Results memvonisnya Not Qualified, seluruhnya karena dokumen wajib yang
    lewat tenggat; skornya sendiri memang di atas batas lulus, jadi angka di sheet
    ini tidak salah — yang salah hanya baris kesimpulannya. Sheet ini juga tidak
    pernah bisa berbunyi PENDING, karena status itu mustahil disimpulkan tanpa tahu
    ada berapa dokumen yang kurang.

    ``status_note`` ditulis sebagai baris tersendiri sebelum "Hasil", satu pola
    dengan catatan fraud dan badword di bawah: setiap sebab gugur yang TIDAK terbaca
    dari angka harus punya barisnya sendiri, supaya "TIDAK LULUS" dengan skor di atas
    batas lulus tidak terbaca seperti salah hitung."""
    if not result_json:
        return
    evaluation = result_json.get("evaluation") or {}
    # phase2 = skor maksimal dikurangi bobot tiap item BELUM_SESUAI.
    phase2 = _scorecard_score(evaluation)
    verif = _to_num(evaluation.get("ai_score_verification"))
    critical = _to_num(evaluation.get("ai_score_critical_compliance_check"))
    # phase3 = phase2 + verifikasi data + pelanggaran kritis.
    if phase2 is not None or verif is not None or critical is not None:
        phase3 = _phase3_score(evaluation)
        phase3 = int(phase3) if phase3 == int(phase3) else phase3
    else:
        phase3 = None
    passing = _to_num(evaluation.get("passing_grade"))
    # max_score diturunkan dari breakdown produk yang diminati (bagian 1),
    # fallback ke maximum_score saat tidak ada produk yang diminati.
    breakdown = []
    if (evaluation.get("cashline_interest") or {}).get("status") == "INTERESTED":
        breakdown.append(("Mega Cashline", 108.75))
    if (evaluation.get("mus_interest") or {}).get("status") == "INTERESTED":
        breakdown.append(("Mega Ultima Shield", 41.25))
    if breakdown:
        max_score = sum(s for _, s in breakdown)
        max_score = int(max_score) if max_score == int(max_score) else max_score
    else:
        max_score = _to_num(evaluation.get("maximum_score"))
    blank = ""
    # 1) Skor Maksimal — sum of interested products (breakdown dihitung di atas).
    sheet.append([
        "Skor Maksimal panggilan ini",
        blank,
        blank if breakdown else (max_score if max_score is not None else blank),
    ])
    for name, score in breakdown:
        sheet.append([f"Tertarik {name}", score, blank])
    if breakdown:
        total = sum(s for _, s in breakdown)
        total = int(total) if total == int(total) else total
        sheet.append(["Skor Maksimal", blank, total])
    # 2) Pengurangan — penilaian scorecard.
    belum = [it for it in (evaluation.get("scorecard_result") or []) if it.get("status") == "BELUM_SESUAI"]
    sheet.append(["Pengurangan – Penilaian scorecard", blank, blank])
    for it in belum:
        w = _to_num(it.get("weight"))
        change = -w if w is not None else blank
        desc = f'{category_label(it.get("category")) or "—"} - {it.get("item_code") or "—"} - {it.get("requirement") or "—"}'
        sheet.append([desc, change, blank])
    # ZERO-SCORE RULE diberi barisnya sendiri: tanpa ini kolom "Perubahan" tidak
    # menjumlah ke "Hasil" (skor maksimal 108,75, pengurangan scorecard hanya -7,5,
    # tetapi hasilnya 0), sehingga terbaca seperti salah hitung.
    if _no_product_interest(evaluation) and max_score is not None:
        belum_weight = sum(
            w for it in belum if (w := _to_num(it.get("weight"))) is not None
        )
        cut = -(max_score - belum_weight)
        sheet.append([
            "Nasabah tidak berminat — tidak tertarik Mega Cashline maupun Mega Ultima "
            "Shield, sehingga skor dipaksa menjadi 0",
            int(cut) if cut == int(cut) else cut,
            blank,
        ])
    delta = None
    if phase2 is not None and max_score is not None:
        delta = phase2 - max_score
        delta = int(delta) if delta == int(delta) else delta
    sheet.append([
        "Skor awal (penilaian scorecard)",
        delta if delta is not None else blank,
        phase2 if phase2 is not None else blank,
    ])
    # 3) Pengurangan — verifikasi data.
    codes = [c for c in (evaluation.get("error_codes") or []) if c.get("error_code") in ("B02", "B03", "B05")]
    used: set = set()
    deductions = []
    # When the dynamic 2-match rule is met the leftover MISMATCH dynamic fields
    # score 0 as a group, so drop them from the deduction breakdown too.
    ch_two_match = card_holder_two_match_satisfied(evaluation.get("card_holder_verification"))
    for v in (evaluation.get("card_holder_verification") or []):
        s = _to_num(v.get("item_score"))
        if s is not None and s < 0:
            if ch_two_match and v.get("field") in CARD_HOLDER_DYNAMIC_FIELDS:
                continue
            deductions.append(("B17", v, s))
    for v in (evaluation.get("cashline_data_verification") or []):
        s = _to_num(v.get("item_score"))
        if s is not None and s < 0:
            deductions.append((code_for_cashline_field(v, codes, used), v, s))
    sheet.append(["Pengurangan – Verifikasi data", blank, blank])
    for code, v, s in deductions:
        field = str(v.get("field") or "").replace("_", " ").title()
        reason = v.get("reason")
        desc = f"{code} - {field}" + (f" - {reason}" if reason else "")
        sheet.append([desc, s, blank])
    after_verif = None
    if phase2 is not None or verif is not None:
        after_verif = (phase2 or 0) + (verif or 0)
        after_verif = int(after_verif) if after_verif == int(after_verif) else after_verif
    sheet.append([
        "Setelah verifikasi data",
        verif if verif is not None else blank,
        after_verif if after_verif is not None else blank,
    ])
    # 4) Pengurangan — pelanggaran kritis.
    sheet.append(["Pengurangan – Pelanggaran kritis", critical if critical is not None else blank, blank])
    # 5) Skor Akhir, Batas Lulus, Hasil.
    sheet.append(["Skor Akhir", blank, phase3 if phase3 is not None else blank])
    bl_label = f"Batas Lulus (90% × {max_score})" if max_score is not None else "Batas Lulus (90% × Skor Maksimal)"
    sheet.append([bl_label, blank, passing if passing is not None else blank])
    # Item non-tolerable yang masih BELUM_SESUAI menggugurkan tiket berapa pun
    # skornya. Bobotnya memang sudah muncul sebagai pengurangan di bagian scorecard
    # di atas, tetapi pengurangan itu saja tidak menjelaskan apa-apa: yang membuat
    # gugur bukan angkanya (skornya bisa tetap di atas batas lulus) melainkan sifat
    # TIDAK DAPAT DITOLERANSI-nya. Tanpa baris ini, ekspor tiket seperti itu terbaca
    # persis seperti salah hitung — 148 dari batas 135, tetapi "TIDAK LULUS".
    for _reason in _non_tolerable_reasons(evaluation):
        sheet.append([f"Tidak dapat ditoleransi — {_reason}", blank, blank])
    # Badword TIDAK mengurangi skor — ia hanya menggugurkan hasilnya. Karena itu
    # skornya bisa berada di ATAS batas lulus sementara "Hasil" berbunyi TIDAK LULUS;
    # tanpa baris ini pembacanya wajar mengira ekspor salah hitung. Kegagalan
    # konsistensi verifikasi statik tidak lagi ditulis di sini, sama dengan kolom AI
    # Status di daftar Results (kebijakan 21 Agustus 2026) — sebabnya terbaca dari
    # item kritikal SC_CL_23_1/23_2 yang gagal.
    # Tiap ucapan badword disebut lengkap dengan timestamp + kutipannya, sama dengan
    # tabel Badword Summary di dashboard.
    badwords = badword_rows(evaluation)
    if badwords:
        sheet.append([badword_fail_reason(evaluation), blank, blank])
        for b in badwords:
            desc = f'{b["evidence"]}' + (f' — {b["reason"]}' if b["reason"] else "")
            sheet.append([desc, blank, blank])
    # Sebab gugur yang datang dari luar evaluasi (dokumen, vonis human). Sama seperti
    # fraud & badword di atas: tidak mengurangi skor, hanya menggugurkan hasilnya.
    if status_note:
        sheet.append([status_note, blank, blank])
    sheet.append(["Hasil", blank, _HASIL_LABELS.get(ai_status, "TIDAK LULUS")])


@router.get("/export_result_xlsx/{result_id}")
def export_result_xlsx(
    result_id: str,
    db: Session = Depends(get_db),
    current_user=Depends(get_evaluation_detail_user),
):
    """Export scorecard rows (one per ``scorecard_result`` item) for a single
    ``done`` result as an XLSX download.

    Capability-nya sama dengan detail penilaian di layar (RESULTS_EVALUATION_DETAIL):
    isi berkasnya persis itu. Tombolnya di Results memang sudah disembunyikan dari
    role tanpa capability tersebut (``isSimpleViewer`` di ResultsView.vue), tetapi
    endpoint-nya dulu tidak ikut menutup — sisi sales tetap bisa menariknya langsung."""
    result = crud.get_result(db, result_id)
    if result is None:
        raise HTTPException(status_code=404, detail="result tidak ditemukan")
    # Berkasnya memuat ringkasan penilaian + scorecard + errorcard satu tiket penuh,
    # jadi cakupannya harus sama dengan halaman detail — bukan sekadar "sudah login".
    ensure_can_view_result(db, current_user, result)
    if result.status != "done":
        raise HTTPException(status_code=422, detail="result belum selesai (done)")
    data = crud.get_result_data(db, result_id)
    workbook = Workbook()
    ringkasan_sheet = workbook.active
    ringkasan_sheet.title = "ringkasan_penilaian_ai"
    ringkasan_sheet.append(_RINGKASAN_COLUMNS)
    scorecard_sheet = workbook.create_sheet("scorecard")
    scorecard_sheet.append(_EXPORT_COLUMNS)
    errorcard_sheet = workbook.create_sheet("errorcard")
    errorcard_sheet.append(_ERROR_COLUMNS)
    cust_id = _customer_id_from_files(result.source_files) or result_id
    if data and data.result_json:
        result_json = data.result_json
        # Apply approved Error Code appeals so the export matches the dashboard
        # (appealed error rows removed, linked scorecard item -> SESUAI, score up).
        appeals = crud.error_code_appeals_for_result(db, result_id)
        evaluation = result_json.get("evaluation") if isinstance(result_json, dict) else None
        if isinstance(evaluation, dict):
            _approved = approved_appeals_only(appeals)
            _flip = [a for a in appeals_that_flip(_approved) if _appeal_kind(a) != "add"]
            evaluation = normalize_static_verification(evaluation)
            # Baris verifikasi dinamis tanpa dua sisi pembanding (Ascend/transkrip
            # kosong) turun ke SKIPPED_NULL, bukan MISMATCH — tidak menerbitkan B17.
            evaluation = normalize_dynamic_verification(evaluation)
            _ds = document_status_map(db, [result]).get(result_id)
            if _ds is not None:
                evaluation = apply_static_document_status(evaluation, _ds[0], _ds[1])
                # Sisi cashline menyusul: baris yang menunggu cover buku tabungan
                # menjadi PENDING selama tenggat H+2.
                evaluation = apply_cashline_document_status(evaluation, _ds[0], _ds[1])
            evaluation = apply_approved_appeals(evaluation, _flip)
            evaluation = apply_approved_card_holder_appeals(evaluation, _flip)
            evaluation = apply_approved_cashline_appeals(evaluation, _flip)
            evaluation = apply_approved_critical_compliance_appeals(evaluation, _flip)
            # 'add' bandings lower the score (attach a new error).
            evaluation = apply_added_score_appeals(evaluation, added_appeals_only(appeals))
            result_json = {**result_json, "evaluation": evaluation}
        # Vonis kanonik — helper yang SAMA dengan kolom AI Status di daftar Results,
        # termasuk aturan dokumen H+2 dan Manual Status yang sudah di-approve. Baris
        # "Hasil" dulu disimpulkan dari dict evaluasi saja dan karenanya buta terhadap
        # keduanya; lihat ``_append_ringkasan_rows``.
        _status = ai_status_for_result(db, result)
        _notes = []
        _gap = data_gap_map(db, [result]).get(result_id)
        if _gap and _status in ("PENDING", "FAIL"):
            _notes.append(data_gap_reasons(_gap))
        if _missing_docs_map(db, [result], {result_id: data.result_json}).get(result_id):
            if _status == "PENDING":
                _sla = " (SLA H+2)" if doc_sla_enabled() else ""
                _notes.append(f"Menunggu dokumen pendukung{_sla}")
            elif _status == "FAIL":
                _notes.append("Dokumen pendukung tidak diunggah sampai tenggat (SLA H+2)")
        _note = " · ".join(_notes) or None
        _append_ringkasan_rows(ringkasan_sheet, result_json, _status, _note)
        _append_scorecard_rows(scorecard_sheet, result_json, cust_id)
        _append_errorcard_rows(errorcard_sheet, result_json, cust_id, appeals)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    # Filename = "<ID>_<timestamp>.xlsx" where the timestamp is the moment the
    # XLSX is generated (e.g. "180310ENv4_20260615103021.xlsx").
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"{cust_id}_{timestamp}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get(
    "/export_verification_xlsx",
    dependencies=[Depends(require(RESULTS_EXPORT_VERIFICATION))],
)
def export_verification_xlsx(
    category: str = Query(..., description="verifikasi_statik | verifikasi_dinamik | cashline_verification | cardholder_verification"),
    campaign: Optional[str] = Query(None, description="Batasi ke satu campaign"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Export agregat SATU kategori verifikasi: semua baris yang tidak cocok pada
    tiket Not Qualified & Pending, dalam satu XLSX.

    Kolomnya menyesuaikan kategori (mis. Cashline Verification membawa "Ketentuan
    Produk" yang tidak ada pada card holder) — lihat ``compute_verification_export``.

    Dibatasi ke campaign yang menjadi cakupan pemanggil: export adalah salinan data
    yang dibawa keluar sistem, jadi ia tidak boleh lebih longgar daripada tabel
    Results yang sudah dibatasi.
    """
    from qc_core.compliance.stats_aggregate import (
        VERIFICATION_EXPORT_CATEGORIES,
        compute_verification_export,
    )
    if category not in VERIFICATION_EXPORT_CATEGORIES:
        raise HTTPException(status_code=422, detail="Kategori export tidak dikenal")
    data = compute_verification_export(db, category, campaign,
                                       effective_campaigns_for(db, current_user))

    workbook = Workbook()
    sheet = workbook.active
    # Judul sheet dibatasi 31 karakter oleh format XLSX; label kategori masih jauh
    # di bawahnya, jadi dipakai apa adanya.
    sheet.title = data["label"][:31]
    sheet.append([title for _key, title in data["columns"]])
    for row in data["rows"]:
        cells = [row.get(key) for key, _title in data["columns"]]
        sheet.append(cells)
        # Kolom "Transkrip" bisa memuat beberapa penyebutan yang dipisah baris baru;
        # tanpa wrap_text, Excel menampilkannya sebagai satu baris berantakan.
        for idx, (key, _title) in enumerate(data["columns"], start=1):
            if key == "transkrip" and isinstance(cells[idx - 1], str) and "\n" in cells[idx - 1]:
                sheet.cell(row=sheet.max_row, column=idx).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"{category}_{timestamp}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get(
    "/export_tickets_xlsx",
    dependencies=[Depends(require(RESULTS_EXPORT_TICKETS))],
)
def export_tickets_xlsx(
    campaign: Optional[str] = Query(None),
    ai_status: Optional[str] = Query(None, description="PASS | FAIL | PENDING"),
    manual_status: Optional[str] = Query(None, description="PASS | FAIL | PENDING"),
    ticket_id: Optional[str] = Query(None),
    am_nip: Optional[str] = Query(None),
    tl_nip: Optional[str] = Query(None),
    agent_nip: Optional[str] = Query(None),
    date_start: Optional[str] = Query(None, description="Batas bawah tanggal (YYYY-MM-DD)"),
    date_end: Optional[str] = Query(None, description="Batas atas tanggal (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Export SEMUA tiket pada rentang tanggal & filter yang sedang dipilih di
    halaman Results — satu baris per tiket, apa pun statusnya.

    Berbeda dari ``/export_verification_xlsx`` yang hanya membawa baris MISMATCH
    pada satu kategori verifikasi untuk tiket Not Qualified & Pending. Yang ini
    tarikan periode: yang Qualified pun ikut, karena pertanyaannya "apa saja yang
    masuk bulan ini", bukan "apa yang harus ditindaklanjuti".

    Parameter penyaringnya sengaja sama dengan ``/results`` dan diproses dengan
    helper yang sama (cakupan role, batas campaign, filter hierarki AM/TL/TLO),
    supaya isi file cocok dengan yang terlihat di layar. Yang TIDAK ditiru:
    paginasi — export selalu mengambil seluruh rentang.
    """
    from qc_core.compliance.stats_aggregate import compute_ticket_export

    d_start = _parse_ymd(date_start)
    d_end = _parse_ymd(date_end)
    role_campaigns = effective_campaigns_for(db, current_user)
    if role_campaigns is not None and campaign:
        allowed = {c.strip().casefold() for c in role_campaigns}
        if campaign.strip().casefold() not in allowed:
            campaign = None
            role_campaigns = []  # di luar cakupan -> hasil kosong, bukan melebar
    scoped_cids = _scoped_customer_ids(db, current_user)
    filter_uids = agent_ids_for_hierarchy_filter(db, am_nip, tl_nip, agent_nip,
                                                 role_campaigns)
    if filter_uids is not None:
        filter_cids = crud.customer_ids_for_agent_ids(db, list(filter_uids))
        scoped_cids = (
            list(set(scoped_cids) & set(filter_cids)) if scoped_cids is not None
            else list(filter_cids)
        )
    _iso = ({"uploaded_by_role": "qc_support"}
            if data_scope_for(db, current_user) == SCOPE_QC_SUPPORT_OWN
            else {"exclude_uploaded_by_role": "qc_support"})

    results, _total = crud.list_results(
        db, campaign=campaign, campaigns=role_campaigns, ticket_id=ticket_id,
        page=1, limit=1_000_000, customer_ids=scoped_cids,
        date_start=d_start, date_end=d_end, **_iso,
    )

    # Filter AI / Manual Status diterapkan di Python: keduanya DITURUNKAN per hasil
    # (tidak tersimpan di baris), persis seperti di /results.
    ai_filter = ai_status.strip().upper() if isinstance(ai_status, str) else None
    manual_filter = manual_status.strip().upper() if isinstance(manual_status, str) else None
    if ai_filter in ("PASS", "FAIL", "PENDING") or manual_filter in ("PASS", "FAIL", "PENDING"):
        rows_done = [r for r in results if r.status == "done"]
        qcs = crud.qc_status_requests_for(db, [str(r.id) for r in rows_done])
        # Helper kanonik yang sama dipakai filter /list_results — lihat catatan di sana.
        ai_map = ai_status_map(db, rows_done)
        kept = []
        for r in rows_done:
            rid = str(r.id)
            ai = ai_map.get(rid)
            if ai_filter and ai != ai_filter:
                continue
            if manual_filter and effective_manual_status(qcs.get(rid), ai) != manual_filter:
                continue
            kept.append(r)
        results = kept

    data = compute_ticket_export(db, results)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tiket"
    sheet.append([title for _key, title in data["columns"]])
    for row in data["rows"]:
        cells = [row.get(key) for key, _title in data["columns"]]
        sheet.append(cells)
        # Details Error memuat satu baris per error code; tanpa wrap_text Excel
        # menampilkannya berdempet jadi satu baris panjang.
        for idx, (key, _title) in enumerate(data["columns"], start=1):
            if key == "details_error" and isinstance(cells[idx - 1], str) and "\n" in cells[idx - 1]:
                sheet.cell(row=sheet.max_row, column=idx).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    span = f"{date_start or 'awal'}_{date_end or 'akhir'}"
    filename = f"tickets_{span}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get(
    "/get_nama_ibu_kandung",
    response_model=NamaIbuKandungResponse,
    dependencies=[Depends(require(RESULTS_EXPORT_VERIFICATION))],
)
def get_nama_ibu_kandung(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Hasil verifikasi statik **nama ibu kandung** SEMUA tiket, tanpa parameter.

    Satu baris per tiket ``done`` yang punya baris verifikasi ``nama_ibu_kandung`` —
    MATCH maupun MISMATCH, apa pun AI Status-nya. Bandingkan dengan
    ``/export_verification_xlsx?category=verifikasi_statik`` yang hanya membawa baris
    MISMATCH pada tiket Not Qualified & Pending.

    Tiap baris berisi ``ticket_id`` (id + timestamp file PDF), ``submit_time`` (dari
    ``tms_cashline``), ``ascend`` (acuan bank), ``transkrip`` (penyebutan nasabah),
    ``match``, ``evidence`` (kutipan pada scorecard SC_CL_23_2), ``similarity``, dan
    ``reason``. Similarity-nya sudah dihitung ulang di Python dan banding yang
    disetujui sudah diterapkan — sama dengan yang tampil di dashboard.

    Dibatasi ke campaign yang menjadi cakupan pemanggil, sama seperti
    ``/export_verification_xlsx``.
    """
    from qc_core.compliance.stats_aggregate import compute_nama_ibu_kandung_rows

    rows = compute_nama_ibu_kandung_rows(db, effective_campaigns_for(db, current_user))
    return {"total": len(rows), "rows": rows}


@router.delete(
    "/delete_ticket",
    response_model=TicketDeleteResponse,
    dependencies=[Depends(require(ADMIN_TICKET_DELETE))],
)
def delete_ticket(
    ticket_id: str = Query(..., description="Ticket id (prefix sebelum '_' pada nama file) yang semua entry-nya akan dihapus"),
    db: Session = Depends(get_db),
):
    """Delete ALL result entries for a ticket (Admin only).

    Removes every Result whose ticket id matches ``ticket_id`` along with its
    dependent rows (result_data, documents, qc_status_requests,
    error_code_appeals) via ON DELETE CASCADE. Returns 404 when no entry matches.
    """
    tid = (ticket_id or "").strip()
    if not tid:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="ticket_id wajib diisi")
    deleted = crud.delete_results_by_ticket_id(db, tid)
    if not deleted:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Tidak ada entry untuk ticket '{tid}'",
        )
    return TicketDeleteResponse(ticket_id=tid, deleted=deleted)
